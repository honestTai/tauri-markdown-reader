"""Markdown → Word (.docx) 导出器。

使用 python-docx 将 Markdown 文本转换为 .docx 文档。
处理能力：
- 标题 H1-H6
- 粗体、斜体、行内代码
- 无序/有序列表
- 表格
- 代码块（等宽字体）
- 引用块
- 图片引用
- LaTeX 公式（转为纯文本占位）
"""

from __future__ import annotations

import os
import re

try:
    from docx import Document
    from docx.shared import Inches, Pt, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml.ns import qn
except ImportError:
    raise ImportError("需要 python-docx 库。安装命令: pip install python-docx")


def _set_cell_shading(cell, color: str):
    """设置单元格背景色。"""
    shading = cell._element.get_or_add_tcPr()
    shading_elm = shading.makeelement(qn("w:shd"), {
        qn("w:fill"): color,
        qn("w:val"): "clear",
    })
    shading.append(shading_elm)


def _add_formatted_paragraph(doc: Document, text: str, style: str | None = None) -> None:
    """添加带内联格式的段落。

    支持的 Markdown 内联格式：**粗体**、*斜体*、`代码`。
    """
    p = doc.add_paragraph(style=style)

    # 解析内联格式：按顺序匹配 **粗体**、*斜体*、`代码`、普通文本
    pattern = re.compile(
        r"(\*\*(.+?)\*\*)|"       # **粗体**
        r"(?<!\*)\*(.+?)\*(?!\*)|" # *斜体*（单星号）
        r"`(.+?)`|"               # `代码`
        r"__(.+?)__|"             # __粗体__
        r"(?<!_)_(.+?)_(?!_)"      # _斜体_
    )

    last_end = 0
    for match in pattern.finditer(text):
        # 添加匹配前的普通文本
        prefix = text[last_end:match.start()]
        if prefix:
            run = p.add_run(prefix)

        # 粗体 **...** 或 __...__
        if match.group(2):
            run = p.add_run(match.group(2))
            run.bold = True
        elif match.group(5):
            run = p.add_run(match.group(5))
            run.bold = True
        # 斜体 *...* 或 _..._
        elif match.group(3):
            run = p.add_run(match.group(3))
            run.italic = True
        elif match.group(6):
            run = p.add_run(match.group(6))
            run.italic = True
        # 行内代码 `...`
        elif match.group(4):
            run = p.add_run(match.group(4))
            run.font.name = "Consolas"
            run.font.size = Pt(10)

        last_end = match.end()

    # 剩余文本
    suffix = text[last_end:]
    if suffix:
        p.add_run(suffix)


def _add_code_block(doc: Document, code: str) -> None:
    """添加代码块（等宽字体段落）。"""
    for line in code.split("\n"):
        p = doc.add_paragraph()
        run = p.add_run(line if line else " ")
        run.font.name = "Consolas"
        run.font.size = Pt(9)
        # 段落间距紧凑
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = Pt(13)


def _add_table(doc: Document, table_text: str) -> None:
    """解析 Markdown 表格文本并创建 Word 表格。"""
    lines = table_text.strip().split("\n")
    if len(lines) < 2:
        return

    # 解析行
    rows_data: list[list[str]] = []
    for line in lines:
        # 跳过分隔行
        if re.match(r"^[\|\s\-:]+$", line):
            continue
        cells = [c.strip() for c in line.strip("|").split("|")]
        rows_data.append(cells)

    if not rows_data:
        return

    max_cols = max(len(r) for r in rows_data)
    table = doc.add_table(rows=len(rows_data), cols=max_cols)
    table.style = "Table Grid"

    for i, row_data in enumerate(rows_data):
        for j, cell_text in enumerate(row_data):
            if j < max_cols:
                cell = table.rows[i].cells[j]
                cell.text = cell_text
                # 表头行加粗
                if i == 0:
                    for paragraph in cell.paragraphs:
                        for run in paragraph.runs:
                            run.bold = True
                    _set_cell_shading(cell, "D9E2F3")


def _add_image(doc: Document, image_path: str, alt_text: str = "") -> None:
    """添加图片到文档。"""
    if os.path.isfile(image_path):
        try:
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = p.add_run()
            run.add_picture(image_path, width=Inches(5.5))
            if alt_text:
                doc.add_paragraph(alt_text).alignment = WD_ALIGN_PARAGRAPH.CENTER
        except Exception:
            # 图片无法加载时，添加文字占位
            doc.add_paragraph(f"[图片: {alt_text or image_path}]")


def convert_markdown_to_docx(
    markdown: str,
    output_path: str,
    *,
    image_base_dir: str | None = None,
) -> str:
    """将 Markdown 文本转换为 .docx 文件。

    Args:
        markdown: Markdown 文本。
        output_path: 输出 .docx 文件路径。
        image_base_dir: 图片文件的基础目录（用于解析相对路径）。

    Returns:
        输出文件路径。
    """
    doc = Document()

    # 设置默认字体
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)

    # 按行解析 Markdown
    lines = markdown.split("\n")
    i = 0
    while i < len(lines):
        line = lines[i]

        # 空行
        if not line.strip():
            i += 1
            continue

        # 标题 H1-H6
        heading_match = re.match(r"^(#{1,6})\s+(.+)$", line)
        if heading_match:
            level = len(heading_match.group(1))
            text = heading_match.group(2).strip()
            _add_formatted_paragraph(doc, text, style=f"Heading {level}")
            i += 1
            continue

        # 水平线
        if re.match(r"^[\-\*_]{3,}$", line.strip()):
            doc.add_paragraph("─" * 60)
            i += 1
            continue

        # 无序列表
        bullet_match = re.match(r"^[\-\*\+]\s+(.+)$", line)
        if bullet_match:
            text = bullet_match.group(1)
            p = doc.add_paragraph(style="List Bullet")
            # 手动添加内联格式
            p.clear()
            _add_run_with_format(p, text)
            i += 1
            continue

        # 有序列表
        num_match = re.match(r"^\d+\.\s+(.+)$", line)
        if num_match:
            text = num_match.group(1)
            p = doc.add_paragraph(style="List Number")
            p.clear()
            _add_run_with_format(p, text)
            i += 1
            continue

        # 引用块
        quote_match = re.match(r"^>\s?(.+)$", line)
        if quote_match:
            text = quote_match.group(1)
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.5)
            run = p.add_run(text)
            run.italic = True
            run.font.color.rgb = RGBColor(100, 100, 100)
            i += 1
            continue

        # 代码块开始 ```
        if line.strip().startswith("```"):
            code_lines: list[str] = []
            i += 1
            while i < len(lines) and not lines[i].strip().startswith("```"):
                code_lines.append(lines[i])
                i += 1
            i += 1  # 跳过结束 ```
            _add_code_block(doc, "\n".join(code_lines))
            continue

        # 表格（以 | 开头）
        if line.strip().startswith("|") and line.strip().endswith("|"):
            table_lines: list[str] = [line]
            i += 1
            while i < len(lines) and lines[i].strip().startswith("|"):
                table_lines.append(lines[i])
                i += 1
            _add_table(doc, "\n".join(table_lines))
            continue

        # 图片 ![alt](path)
        image_match = re.match(r"^!\[(.*?)\]\((.*?)\)$", line.strip())
        if image_match:
            alt = image_match.group(1)
            img_path = image_match.group(2)
            # 如果是相对路径，拼基础目录
            if image_base_dir and not os.path.isabs(img_path):
                img_path = os.path.join(image_base_dir, img_path)
            _add_image(doc, img_path, alt)
            i += 1
            continue

        # LaTeX 公式（占位处理，Part 3 后续可渲染为图片）
        if line.strip().startswith("$$"):
            formula_lines: list[str] = []
            i += 1
            while i < len(lines) and not lines[i].strip().startswith("$$"):
                formula_lines.append(lines[i])
                i += 1
            i += 1  # 跳过结束 $$
            formula_text = "\n".join(formula_lines).strip()
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = p.add_run(f"[公式: {formula_text[:100]}]")
            run.italic = True
            run.font.color.rgb = RGBColor(128, 128, 128)
            continue

        # 普通段落
        _add_formatted_paragraph(doc, line)
        i += 1

    doc.save(output_path)
    return output_path


def _add_run_with_format(paragraph, text: str) -> None:
    """辅助函数：向段落添加带内联格式的 run。"""
    pattern = re.compile(
        r"(\*\*(.+?)\*\*)|"
        r"(?<!\*)\*(.+?)\*(?!\*)|"
        r"`(.+?)`|"
        r"__(.+?)__|"
        r"(?<!_)_(.+?)_(?!_)"
    )

    last_end = 0
    for match in pattern.finditer(text):
        prefix = text[last_end:match.start()]
        if prefix:
            paragraph.add_run(prefix)

        if match.group(2):
            run = paragraph.add_run(match.group(2))
            run.bold = True
        elif match.group(5):
            run = paragraph.add_run(match.group(5))
            run.bold = True
        elif match.group(3):
            run = paragraph.add_run(match.group(3))
            run.italic = True
        elif match.group(6):
            run = paragraph.add_run(match.group(6))
            run.italic = True
        elif match.group(4):
            run = paragraph.add_run(match.group(4))
            run.font.name = "Consolas"

        last_end = match.end()

    suffix = text[last_end:]
    if suffix:
        paragraph.add_run(suffix)
