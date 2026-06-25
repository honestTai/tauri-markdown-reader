"""docx_parser 单元测试。"""

import os
import sys
import tempfile
from io import BytesIO

# 将父目录加入搜索路径
sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from parsers.docx_parser import (
    convert_docx_to_markdown,
    _table_to_markdown,
    _detect_heading_level,
    _is_list_item,
)


def _create_test_docx() -> str | None:
    """创建一个包含标题、段落和表格的最小测试 .docx。"""
    try:
        from docx import Document
    except ImportError:
        return None

    doc = Document()

    # 一级标题
    doc.add_heading("Test Document", level=1)

    # 带格式的段落
    p = doc.add_paragraph()
    r1 = p.add_run("bold text")
    r1.bold = True
    p.add_run(" and ")
    r2 = p.add_run("italic text")
    r2.italic = True
    p.add_run(".")

    # 二级标题
    doc.add_heading("Section One", level=2)

    # 普通段落
    doc.add_paragraph("This is a normal paragraph with some content.")

    # 无序列表
    doc.add_paragraph("Bullet item 1", style="List Bullet")
    doc.add_paragraph("Bullet item 2", style="List Bullet")

    # 表格
    table = doc.add_table(rows=3, cols=3)
    table.style = "Table Grid"
    header_cells = table.rows[0].cells
    header_cells[0].text = "Name"
    header_cells[1].text = "Age"
    header_cells[2].text = "City"
    for i, (name, age, city) in enumerate(
        [("Alice", "30", "NYC"), ("Bob", "25", "LA")], start=1
    ):
        table.rows[i].cells[0].text = name
        table.rows[i].cells[1].text = age
        table.rows[i].cells[2].text = city

    buf = BytesIO()
    doc.save(buf)
    buf.seek(0)

    # 写入临时文件
    fd, path = tempfile.mkstemp(suffix=".docx")
    with os.fdopen(fd, "wb") as f:
        f.write(buf.read())
    return path


def test_convert_docx_basic():
    """测试基本 docx 转 Markdown。"""
    docx_path = _create_test_docx()
    if docx_path is None:
        print("跳过: python-docx 不可用")
        return

    try:
        result = convert_docx_to_markdown(docx_path)

        # 验证关键内容存在
        assert "# Test Document" in result
        assert "## Section One" in result
        assert "**bold text**" in result
        assert "*italic text*" in result
        assert "- Bullet item 1" in result
        assert "- Bullet item 2" in result
        assert "| Name |" in result
        assert "| Alice |" in result
        assert "| Bob |" in result

        print("通过: test_convert_docx_basic")
    finally:
        os.unlink(docx_path)


def test_empty_table():
    """测试空表格不产生输出。"""
    try:
        from docx import Document
    except ImportError:
        print("跳过: python-docx 不可用")
        return

    doc = Document()
    doc.add_table(rows=0, cols=0)
    assert _table_to_markdown(doc.tables[0]) == ""
    print("通过: test_empty_table")


def test_heading_detection():
    """测试标题级别检测。"""
    try:
        from docx import Document
    except ImportError:
        print("跳过: python-docx 不可用")
        return

    doc = Document()
    h1 = doc.add_heading("一级", level=1)
    assert _detect_heading_level(h1) == 1

    h3 = doc.add_heading("三级", level=3)
    assert _detect_heading_level(h3) == 3

    p = doc.add_paragraph("普通文本")
    assert _detect_heading_level(p) is None

    print("通过: test_heading_detection")


def test_list_detection():
    """测试列表项检测。"""
    try:
        from docx import Document
    except ImportError:
        print("跳过: python-docx 不可用")
        return

    doc = Document()
    bullet = doc.add_paragraph("无序", style="List Bullet")
    assert _is_list_item(bullet) == "-"

    numbered = doc.add_paragraph("有序", style="List Number")
    assert _is_list_item(numbered) == "1."

    normal = doc.add_paragraph("普通")
    assert _is_list_item(normal) is None

    print("通过: test_list_detection")


if __name__ == "__main__":
    test_convert_docx_basic()
    test_empty_table()
    test_heading_detection()
    test_list_detection()
    print("\n所有 docx_parser 测试通过！")
